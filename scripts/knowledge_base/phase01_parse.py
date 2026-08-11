#!/usr/bin/env python3
from __future__ import annotations

import argparse, csv, datetime as dt, hashlib, json, re, shutil, sys
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Iterable
from openpyxl import load_workbook

PHASE="PHASE-01"
PORTALS={"employee":"01_员工端","center":"02_中心管理端","tech":"03_技术后台端"}
RUNTIME={"employee":"employee","center":"center","tech":"admin"}
PAGE_FILES={
"employee":["Knowledge Base/01 完整的页面架构/1-1 员工首页.xlsx","Knowledge Base/01 完整的页面架构/1-2员工全层级页面.xlsx"],
"center":["Knowledge Base/01 完整的页面架构/2-1 中心首页.xlsx","Knowledge Base/01 完整的页面架构/2-2中心全层级页面.xlsx"],
"tech":["Knowledge Base/01 完整的页面架构/3-1技术-首页.xlsx","Knowledge Base/01 完整的页面架构/3-2技术-全层级页面.xlsx"]}
PROC_ROOT=Path("Knowledge Base/02 业务流程 表单 字段/S1 三端业务流程表单字段包")
S0=Path("Knowledge Base/02 业务流程 表单 字段/S0 全部业务流程简表.xlsx")
S1=PROC_ROOT/"00_三端业务流程表单字段总索引.xlsx"
DB=Path("Knowledge Base/03 数据库需求规则")
DBFILES={
"data_sources":DB/"02_数据字典/01_数据源规划.csv","schemas":DB/"02_数据字典/02_Schema分域.csv",
"tables":DB/"02_数据字典/03_全量表清单.csv","fields":DB/"02_数据字典/04_全量字段字典.csv",
"relations":DB/"02_数据字典/05_主外键关系.csv","indexes":DB/"02_数据字典/06_索引设计.csv",
"mapping":DB/"02_数据字典/07_流程落表映射.csv","process_catalog":DB/"04_初始化数据/01_process_catalog.csv",
"state_catalog":DB/"04_初始化数据/02_state_catalog.csv","rule_catalog":DB/"04_初始化数据/03_business_rule_catalog.csv",
"interface_catalog":DB/"04_初始化数据/04_interface_catalog.csv","linkage_catalog":DB/"04_初始化数据/05_three_endpoint_linkage.csv",
"source_map":DB/"07_源材料映射/01_原始字段到目标数据库全量映射.csv",
"unique_map":DB/"07_源材料映射/02_唯一字段规范化映射.csv"}
EXPECTED={f"P{i:03d}" for i in range(1,127)}
PCODE=re.compile(r"\bP?(00[1-9]|0[1-9]\d|1[01]\d|12[0-6])\b",re.I)
PFILE=re.compile(r"(\d{3})_(.+)\.xlsx$",re.I)
SECTIONS={"overview":["流程总览","流程概览"],"forms":["表单清单","表单目录"],
"fields":["字段字典","字段清单"],"states":["状态与审批","状态审批"],
"rules":["规则与接口","接口与规则"],"linkage":["三端联动"]}
HINTS=["一级","二级","三级","页面","流程","字段","表单","状态","规则","接口","权限","模块","名称","编码","范围","敏感","动作","审批"]
PALIASES={"level_1":["一级页面","一级菜单","一级模块","一级板块","一级"],"level_2":["二级页面","二级菜单","二级模块","二级板块","二级"],
"level_3":["三级页面","三级菜单","三级模块","三级板块","三级"],"display":["页面名称","功能页面","功能名称","展示名称","页面"],
"process":["process_code","process_codes","流程码","流程编码","业务流程编码","关联流程"],
"permission":["permission_code","permission_codes","权限码","权限代码","权限"],"scope":["data_scope","数据范围","数据权限范围"],
"sensitive":["sensitive_level","敏感级别","数据敏感级别","敏感等级","安全等级"],
"mobile":["mobile_access","移动端","手机端","移动访问","移动可达"],"status":["status","实施状态","页面状态"],
"aliases":["aliases","别名","历史名称","兼容名称"],"notes":["notes","备注","说明"]}
DBALIASES={"database":["数据库","数据库名","database","database_name","db_name","data_source"],
"schema":["schema","schema_name","模式","模式名"],"table":["table","table_name","表","表名","物理表","物理表名"],
"column":["column","column_name","字段","字段名","列","列名"],"process":["process_code","流程编码","流程码","业务流程编码"],
"main_table":["main_table","主表","权威主表","主表名"],"method":["method","http_method","请求方法"],
"path":["path","api_path","接口路径","url"]}
PERMKW=["权限","可见","可编辑","可导出","角色","数据范围","敏感","审批人","处理人","责任人","动作","操作","授权","mfa","step","回避","四眼","职责分离","scope","permission"]
SECRET=[re.compile(r"github_pat_\w+",re.I),re.compile(r"gh[pousr]_\w+",re.I),re.compile(r"\bsk-[\w-]{12,}\b",re.I)]
PII=[(re.compile(r"^1[3-9]\d{9}$"),"[REDACTED_PHONE]"),(re.compile(r"^\d{17}[\dXx]$"),"[REDACTED_ID]"),
(re.compile(r"^\d{16,19}$"),"[REDACTED_LONG_NUMBER]"),(re.compile(r"^[\w.%+-]+@[\w.-]+\.[A-Za-z]{2,}$"),"[REDACTED_EMAIL]")]

def text(v): return re.sub(r"\s+"," ",str(v).replace("\r","\n")).strip() if v is not None else ""
def key(v): return re.sub(r"[\s_\-/:：()（）\[\]【】]+","",text(v).lower())
def name(v): return re.sub(r"[\s_\-—–·,，。.;；:：/\\()（）\[\]【】]+","",text(v)).lower()
def safe(v):
    if isinstance(v,(dt.datetime,dt.date,dt.time)): return v.isoformat()
    if isinstance(v,float) and v.is_integer(): return int(v)
    return v
def redact(v):
    v=safe(v)
    if not isinstance(v,str): return v
    s=v.strip()
    if any(p.search(s) for p in SECRET): return "[REDACTED_SECRET]"
    for p,r in PII:
        if p.fullmatch(s): return r
    return v
def sha(path):
    h=hashlib.sha256()
    with path.open("rb") as f:
        for b in iter(lambda:f.read(1048576),b""): h.update(b)
    return h.hexdigest()
def jwrite(path,obj):
    path.parent.mkdir(parents=True,exist_ok=True); path.write_text(json.dumps(obj,ensure_ascii=False,indent=2,default=str)+"\n",encoding="utf-8")
def jlwrite(path,rows):
    path.parent.mkdir(parents=True,exist_ok=True)
    with path.open("w",encoding="utf-8",newline="\n") as f:
        for r in rows: f.write(json.dumps(r,ensure_ascii=False,separators=(",",":"),default=str)+"\n")
def headers(vals):
    seen=Counter(); out=[]
    for i,v in enumerate(vals,1):
        b=text(v) or f"column_{i}"; seen[b]+=1; out.append(b if seen[b]==1 else f"{b}__{seen[b]}")
    return out
def header_idx(rows):
    best=(0,-1.0)
    for i,row in enumerate(rows[:30]):
        xs=[text(x) for x in row if text(x)]
        if not xs: continue
        score=len(xs)+3*sum(any(h in x.lower() for h in HINTS) for x in xs)+len(set(xs))/len(xs)
        if score>best[1]: best=(i,score)
    return best[0]
def xlsx(path):
    wb=load_workbook(path,read_only=True,data_only=False); metas=[]; recs=[]
    try:
        for ws in wb.worksheets:
            rr=[tuple(r) for r in ws.iter_rows(values_only=True)]
            ne=[r for r in rr if any(text(v) for v in r)]
            if not ne: metas.append({"sheet":ws.title,"headers":[],"header_row":None,"rows":0}); continue
            hi=header_idx(ne); hv=ne[hi]; hs=headers(hv); n=-1; src_h=1
            for sr,row in enumerate(rr,1):
                if not any(text(v) for v in row): continue
                n+=1
                if n==hi: src_h=sr; break
            count=0
            for sr,row in enumerate(rr,1):
                if sr<=src_h or not any(text(v) for v in row): continue
                row=list(row)+[None]*max(0,len(hs)-len(row))
                recs.append({"sheet":ws.title,"row":sr,"data":{hs[i]:redact(row[i]) for i in range(len(hs))}}); count+=1
            metas.append({"sheet":ws.title,"headers":hs,"header_row":src_h,"rows":count,"max_row":ws.max_row,"max_column":ws.max_column})
    finally: wb.close()
    return metas,recs
def csvread(path):
    raw=path.read_bytes(); txt=enc=None
    for e in ("utf-8-sig","utf-8","gb18030"):
        try: txt=raw.decode(e); enc=e; break
        except UnicodeDecodeError: pass
    if txt is None: raise ValueError("unsupported encoding")
    try: d=csv.Sniffer().sniff(txt[:65536],delimiters=",\t;")
    except csv.Error: d=csv.excel
    rows=list(csv.reader(txt.splitlines(),d))
    if not rows: return [],[],enc
    hi=header_idx(rows); hs=headers(rows[hi]); out=[]
    for sr,row in enumerate(rows[hi+1:],hi+2):
        if not any(text(v) for v in row): continue
        row=list(row)+[""]*max(0,len(hs)-len(row))
        out.append({"row":sr,"data":{hs[i]:redact(row[i]) for i in range(len(hs))}})
    return hs,out,enc
def hmap(hs,aliases):
    nk={key(h):h for h in hs}
    for a in aliases:
        if key(a) in nk: return nk[key(a)]
    for a in aliases:
        for k0,h in nk.items():
            if key(a) and (key(a) in k0 or k0 in key(a)): return h
    return None
def codes(v):
    out=[]
    for m in PCODE.finditer(text(v)):
        c="P"+m.group(1)
        if c not in out: out.append(c)
    return out
def rowcode(data):
    for k0,v in data.items():
        if "流程" in key(k0) or "process" in key(k0):
            cs=codes(v)
            if cs:return cs[0]
    for v in data.values():
        cs=codes(v)
        if cs:return cs[0]
    return None
def section(sheet):
    n=key(sheet)
    for s in ("linkage","rules","states","overview","forms","fields"):
        for p in SECTIONS[s]:
            if key(p) in n:return s
    return None
def slist(v): return [x.strip() for x in re.split(r"[,，;；|/、\n]+",text(v)) if x.strip()]
def sensitive(v):
    s=text(v).upper().replace("_","-")
    for a,b in {"P0":"DATA-L1","P1":"DATA-L2","P2":"DATA-L3","P3":"DATA-L4","L1":"DATA-L1","L2":"DATA-L2","L3":"DATA-L3","L4":"DATA-L4"}.items():
        if a in s:return b
    return s or "UNKNOWN"
def mobile(v):
    s=text(v).lower()
    if not s:return "UNKNOWN"
    if s in {"是","支持","完整","可办理","full"} or "full" in s:return "full"
    if s in {"部分","限制","有限","只读","limited"} or "limited" in s:return "limited"
    if s in {"否","不支持","不可","no"}:return "no"
    return s
def mdtable(cols,rows):
    esc=lambda v:text(v).replace("|","\\|")
    return "\n".join(["| "+" | ".join(map(esc,cols))+" |","| "+" | ".join("---" for _ in cols)+" |"]+
                    ["| "+" | ".join(esc(v) for v in r)+" |" for r in rows])

class Pipeline:
    def __init__(self,root):
        self.root=Path(root).resolve(); self.kb=self.root/"Knowledge Base"; self.impl=self.root/"docs/implementation"
        self.contract=self.impl/"contracts/phase-01"; self.gaps=[]; self.hard=[]; self.manifest=[]
        self.pages=[]; self.processes=[]; self.sections={s:[] for s in SECTIONS}; self.perms=[]; self.db={}
    def gap(self,severity,category,source,description,impact):
        self.gaps.append({"id":f"GAP-{len(self.gaps)+1:04d}","severity":severity,"category":category,
                          "source":source,"description":description,"impact":impact,"handling":"UNKNOWN：记录来源，不擅自修正"})
    def scan(self):
        for p in sorted(x for x in self.kb.rglob("*") if x.is_file()):
            rel=p.relative_to(self.root).as_posix(); item={"path":rel,"size":p.stat().st_size,"sha256":sha(p),"ext":p.suffix.lower(),"status":"PASS"}
            try:
                if p.suffix.lower()==".xlsx":
                    m,_=xlsx(p); item|={"parser":"openpyxl","sheets":m}
                elif p.suffix.lower()==".csv":
                    h,r,e=csvread(p); item|={"parser":"csv","encoding":e,"headers":h,"rows":len(r)}
                elif p.suffix.lower()==".json":
                    o=json.loads(p.read_text(encoding="utf-8-sig")); item|={"parser":"json","root":type(o).__name__,"items":len(o) if isinstance(o,(list,dict)) else None}
                elif p.suffix.lower() in {".md",".txt",".sql",".mmd"}:
                    t=p.read_text(encoding="utf-8-sig"); item|={"parser":"text","lines":t.count("\n")+1}
                else:item["parser"]="metadata"
            except Exception as e:
                item|={"status":"FAIL","error":f"{type(e).__name__}: {e}"}; self.hard.append(f"parse:{rel}"); self.gap("BLOCKER","PARSE",rel,item["error"],"全量 KB 机器解析不完整")
            self.manifest.append(item)
    def parse_pages(self):
        self.page_counts={}
        for portal,files in PAGE_FILES.items():
            for rel in files:
                p=self.root/rel
                if not p.exists(): self.hard.append("page-missing:"+rel); self.gap("BLOCKER","PAGE",rel,"权威页面 Excel 缺失","六份页面 100% 解析失败"); continue
                metas,recs=xlsx(p); by=defaultdict(list)
                for r in recs:by[r["sheet"]].append(r)
                count=0
                for meta in metas:
                    hs=meta["headers"]; mp={k:hmap(hs,a) for k,a in PALIASES.items()}; carry={}
                    for r in by[meta["sheet"]]:
                        d=r["data"]
                        for lv in ("level_1","level_2","level_3"):
                            h=mp[lv]; v=text(d.get(h)) if h else ""
                            if v:carry[lv]=v
                        l1=carry.get("level_1"); l2=carry.get("level_2"); l3=carry.get("level_3")
                        disp=text(d.get(mp["display"])) if mp["display"] else ""; disp=disp or text(l3 or l2 or l1)
                        if not l1 or not disp:continue
                        cs=codes(d.get(mp["process"])) if mp["process"] else []
                        pc=slist(d.get(mp["permission"])) if mp["permission"] else []
                        scope=text(d.get(mp["scope"])) if mp["scope"] else ""; sen=sensitive(d.get(mp["sensitive"])) if mp["sensitive"] else "UNKNOWN"
                        mob=mobile(d.get(mp["mobile"])) if mp["mobile"] else "UNKNOWN"; st=text(d.get(mp["status"])) if mp["status"] else ""
                        st="implemented" if any(x in st.lower() for x in ("implemented","已实现","上线")) else "deprecated" if any(x in st.lower() for x in ("deprecated","废弃","下线")) else "planned"
                        src=f"{Path(rel).name}:{meta['sheet']}:R{r['row']}:{hashlib.sha256(json.dumps(d,ensure_ascii=False,default=str,sort_keys=True).encode()).hexdigest()[:12]}"
                        rec={"portal_code":portal,"runtime_portal":RUNTIME[portal],"source_file":rel,"source_sheet":meta["sheet"],"source_row":r["row"],"source_key":src,
                             "level_1":l1,"level_2":l2 or None,"level_3":l3 or None,"display_name":disp,"aliases":slist(d.get(mp["aliases"])) if mp["aliases"] else [],
                             "route_name":None,"route_path":None,"implementation_path":None,"permission_codes":pc,"process_codes":cs,
                             "data_scope":scope or "UNKNOWN","sensitive_level":sen,"mobile_access":mob,"status":st,"notes":text(d.get(mp["notes"])) if mp["notes"] else None}
                        self.pages.append(rec); count+=1
                        for fld in ("data_scope","sensitive_level","mobile_access"):
                            if rec[fld]=="UNKNOWN":self.gap("WARN","PAGE_FACT",f"{rel}#{meta['sheet']}!R{r['row']}",f"{disp} 缺少可来源化 {fld}","实现前必须补齐，当前保持 UNKNOWN")
                self.page_counts[rel]=count
                if count==0:self.hard.append("page-zero:"+rel)
    def parse_indexes_and_processes(self):
        _,s0r=xlsx(self.root/S0); _,s1r=xlsx(self.root/S1)
        self.s0=[{"process_code":rowcode(r["data"]),"source_file":S0.as_posix(),**r} for r in s0r if rowcode(r["data"])]
        self.s1=[{"process_code":rowcode(r["data"]),"source_file":S1.as_posix(),**r} for r in s1r if rowcode(r["data"])]
        self.books={p:{} for p in PORTALS}; self.bookmeta=[]; self.bookcounts=Counter()
        for portal,dirname in PORTALS.items():
            for p in sorted((self.root/PROC_ROOT/dirname).rglob("*.xlsx")):
                if p.name=="00_目录索引.xlsx":continue
                m=PFILE.match(p.name)
                if not m:continue
                code="P"+m.group(1); rel=p.relative_to(self.root).as_posix()
                if code in self.books[portal]:self.hard.append(f"dup:{portal}:{code}");self.gap("BLOCKER","PROCESS_DUP",rel,f"{portal} {code} 重复工作簿","流程唯一性失败");continue
                self.books[portal][code]=p; self.bookcounts[portal]+=1
                metas,recs=xlsx(p); by=defaultdict(list)
                for r in recs:by[r["sheet"]].append(r)
                cls=defaultdict(list)
                for meta in metas:
                    s=section(meta["sheet"])
                    if s:cls[s].append(meta["sheet"])
                missing=[s for s in SECTIONS if not cls[s]]
                if missing:self.gap("WARN","PROCESS_SECTION",rel,f"{code}/{portal} 未识别 Sheet: {missing}","对应合同保持 UNKNOWN")
                self.bookmeta.append({"process_code":code,"portal_code":portal,"source_file":rel,"name_from_file":m.group(2),"sheets":metas,"sections":dict(cls),"missing":missing})
                for s,sheets in cls.items():
                    for sh in sheets:
                        for r in by[sh]:
                            row={"process_code":code,"portal_code":portal,"source_file":rel,"source_sheet":sh,"source_row":r["row"],"data":r["data"]}
                            self.sections[s].append(row)
                            facts={k:v for k,v in r["data"].items() if text(v) and any(key(w) in key(k) for w in PERMKW)}
                            if facts:self.perms.append({"process_code":code,"portal_code":portal,"source_file":rel,"source_sheet":sh,"source_row":r["row"],"facts":facts})
        for portal in PORTALS:
            miss=sorted(EXPECTED-set(self.books[portal])); extra=sorted(set(self.books[portal])-EXPECTED)
            if miss or extra or len(self.books[portal])!=126:
                self.hard.append(f"coverage:{portal}");self.gap("BLOCKER","PROCESS_COVERAGE",str(PROC_ROOT),f"{portal}: count={len(self.books[portal])}, missing={miss}, extra={extra}","P001-P126 三端覆盖失败")
        self.s0codes=Counter(x["process_code"] for x in self.s0); self.s1codes=Counter(x["process_code"] for x in self.s1)
        for tag,c in (("S0",self.s0codes),("S1_INDEX",self.s1codes)):
            miss=sorted(EXPECTED-set(c));dup=sorted(k for k,v in c.items() if v>1);extra=sorted(set(c)-EXPECTED)
            if miss or dup or extra:self.gap("WARN","PROCESS_INDEX",tag,f"missing={miss}, duplicate={dup}, extra={extra}","不自动修正索引")
    def parse_db(self):
        for k,rel in DBFILES.items():
            p=self.root/rel
            if not p.exists():self.gap("WARN","DB_SOURCE",rel.as_posix(),"数据库来源缺失","对应合同 UNKNOWN");continue
            try:
                h,r,e=csvread(p);self.db[k]={"source_file":rel.as_posix(),"headers":h,"records":r,"encoding":e}
            except Exception as ex:self.hard.append("db:"+k);self.gap("BLOCKER","DB_PARSE",rel.as_posix(),str(ex),"数据库统计失败")
        def dcol(item,canon):return hmap(item["headers"],DBALIASES[canon])
        def cnt(k,canon):
            i=self.db.get(k)
            if not i:return (0,None)
            c=dcol(i,canon); vals={text(r["data"].get(c)) for r in i["records"] if c and text(r["data"].get(c))}
            return (len(vals) if vals else len(i["records"]),c)
        self.dbcount,self.dbcol=cnt("data_sources","database");self.schemacount,self.schemacol=cnt("schemas","schema")
        ti=self.db.get("tables"); ids=set()
        if ti:
            cs=[dcol(ti,x) for x in ("database","schema","table")]
            for r in ti["records"]:
                pp=[text(r["data"].get(c)) if c else "" for c in cs]
                if pp[-1]:ids.add(".".join(x for x in pp if x))
        self.tablecount=len(ids) if ids else len(ti["records"]) if ti else 0
        self.dbfieldrows=len(self.db.get("fields",{}).get("records",[]));self.relrows=len(self.db.get("relations",{}).get("records",[]));self.idxrows=len(self.db.get("indexes",{}).get("records",[]))
        self.pmap=defaultdict(list);mi=self.db.get("mapping")
        if mi:
            pc=dcol(mi,"process");sc=dcol(mi,"schema");tc=dcol(mi,"table") or dcol(mi,"main_table")
            for r in mi["records"]:
                cs=codes(r["data"].get(pc)) if pc else [];c=cs[0] if cs else rowcode(r["data"])
                if c:self.pmap[c].append({"source_file":mi["source_file"],"source_row":r["row"],"schema":text(r["data"].get(sc)) if sc else "UNKNOWN","table":text(r["data"].get(tc)) if tc else "UNKNOWN"})
        missing=sorted(EXPECTED-set(self.pmap))
        if missing:self.gap("WARN","DB_MAPPING",DBFILES["mapping"].as_posix(),f"{len(missing)} processes without mapping: {missing}","权威 Schema/主表保持 UNKNOWN")
        for label,actual,decl in (("database",self.dbcount,3),("schema",self.schemacount,46),("table",self.tablecount,265)):
            if actual and actual!=decl:self.gap("WARN","DB_COUNT",str(DB),f"{label}: declared={decl}, actual={actual}","以当前机器解析统计为事实并保留差异")
        self.ddl=self.ddl_inventory()
    def ddl_inventory(self):
        pats={"database":re.compile(r"\bCREATE\s+DATABASE\s+(?:IF\s+NOT\s+EXISTS\s+)?[\"']?([\w]+)",re.I),
              "schema":re.compile(r"\bCREATE\s+SCHEMA\s+(?:IF\s+NOT\s+EXISTS\s+)?[\"']?([\w]+)",re.I),
              "table":re.compile(r"\bCREATE\s+TABLE\s+(?:IF\s+NOT\s+EXISTS\s+)?([\"\w.]+)",re.I),
              "index":re.compile(r"\bCREATE\s+(?:UNIQUE\s+)?INDEX\s+(?:CONCURRENTLY\s+)?(?:IF\s+NOT\s+EXISTS\s+)?[\"']?([\w]+)",re.I)}
        out={k:[] for k in pats};files=[]
        for p in sorted((self.root/DB/"03_SQL_DDL").rglob("*.sql")):
            rel=p.relative_to(self.root).as_posix();t=p.read_text(encoding="utf-8-sig");files.append({"path":rel,"lines":t.count("\n")+1})
            for k,pat in pats.items():
                for m in pat.finditer(t):out[k].append({"name":m.group(1).replace('"',""),"source_file":rel})
        return {"files":files,"counts":{k:len(v) for k,v in out.items()},"creates":out}
    def process_name(self,records):
        out=[]
        for r in records:
            for k0,v in r["data"].items():
                if "名称" in key(k0) or "流程名" in key(k0) or key(k0) in {"processname","name"}:
                    if text(v) and text(v) not in out:out.append(text(v))
        return out
    def build_process_catalog(self):
        dbpc=defaultdict(list)
        for r in self.db.get("process_catalog",{}).get("records",[]):
            c=rowcode(r["data"])
            if c:dbpc[c].append(r)
        s0=defaultdict(list);s1=defaultdict(list)
        for r in self.s0:s0[r["process_code"]].append(r)
        for r in self.s1:s1[r["process_code"]].append(r)
        for c in sorted(EXPECTED):
            ps={};fns=[]
            for portal in PORTALS:
                p=self.books[portal].get(c)
                if p:
                    rel=p.relative_to(self.root).as_posix();m=PFILE.match(p.name);fns.append(m.group(2) if m else p.stem)
                    bm=next((x for x in self.bookmeta if x["process_code"]==c and x["portal_code"]==portal),None)
                    ps[portal]={"source_file":rel,"name_from_file":m.group(2) if m else p.stem,"sections":bm["sections"] if bm else {},"missing_sections":bm["missing"] if bm else list(SECTIONS)}
            sn={"db":self.process_name(dbpc[c]),"s0":self.process_name(s0[c]),"s1":self.process_name(s1[c]),"files":sorted(set(fns))}
            norms={name(x) for a in sn.values() for x in a if name(x)}
            if len(norms)>1:self.gap("WARN","PROCESS_NAME",c,f"cross-source names={sn}","不自动改名")
            canonical=(sn["db"] or sn["s0"] or sn["files"] or ["UNKNOWN"])[0]
            self.processes.append({"process_code":c,"canonical_name":canonical,"source_names":sn,"portal_sources":ps,"database_mappings":self.pmap.get(c,[])})
    def build_trace(self):
        self.trace=[];forms=defaultdict(list);fields=defaultdict(list);perms=defaultdict(list)
        for r in self.sections["forms"]:forms[(r["process_code"],r["portal_code"])].append(r)
        for r in self.sections["fields"]:fields[(r["process_code"],r["portal_code"])].append(r)
        for r in self.perms:perms[(r["process_code"],r["portal_code"])].append(r)
        ii=self.db.get("interface_catalog",{});interfaces=defaultdict(list)
        for r in ii.get("records",[]):
            c=rowcode(r["data"])
            if c:interfaces[c].append(r)
        self.api=[];mh=hmap(ii.get("headers",[]),DBALIASES["method"]) if ii else None;ph=hmap(ii.get("headers",[]),DBALIASES["path"]) if ii else None
        if ii:
            for r in ii["records"]:
                m=text(r["data"].get(mh)) if mh else "";p=text(r["data"].get(ph)) if ph else ""
                if m or p:self.api.append({"process_code":rowcode(r["data"]) or "UNKNOWN","source_file":ii["source_file"],"source_row":r["row"],"method":m or "UNKNOWN","path":p or "UNKNOWN"})
        missing=0
        pidx={x["process_code"]:x for x in self.processes}
        for pg in self.pages:
            cs=pg["process_codes"] or ["UNKNOWN"];missing+=not bool(pg["process_codes"])
            for c in cs:
                k=(c,pg["portal_code"]);pr=pidx.get(c)
                self.trace.append({"page":{"portal_code":pg["portal_code"],"display_name":pg["display_name"],"source_file":pg["source_file"],"source_sheet":pg["source_sheet"],"source_row":pg["source_row"],"source_key":pg["source_key"]},
                "process_code":c,"process_sources":pr["portal_sources"] if pr else {},
                "forms":{"count":len(forms[k]),"source_file":forms[k][0]["source_file"] if forms[k] else "UNKNOWN"},
                "fields":{"count":len(fields[k]),"source_file":fields[k][0]["source_file"] if fields[k] else "UNKNOWN"},
                "database":self.pmap.get(c,[]) if c!="UNKNOWN" else [],
                "permissions":{"count":len(perms[k]),"source_file":perms[k][0]["source_file"] if perms[k] else "UNKNOWN"},
                "api_or_interface":{"count":len(interfaces.get(c,[])),"source_file":ii.get("source_file","UNKNOWN") if interfaces.get(c) else "UNKNOWN"}})
        if missing:self.gap("WARN","PAGE_PROCESS","六份页面 Excel",f"{missing} page records without sourced process_code","page→process 保持 UNKNOWN")
    def summary(self):
        self.formrows=len(self.sections["forms"]);self.fieldrows=len(self.sections["fields"])
        checks={"six_page_excels":len(getattr(self,"page_counts",{}))==6 and all(self.page_counts.values()),
        "page_source_trace":bool(self.pages) and all(p["source_file"] and p["source_sheet"] and p["source_key"] for p in self.pages),
        "process_catalog_126":len(self.processes)==126,
        "process_workbooks_378":sum(self.bookcounts.values())==378 and all(self.bookcounts[p]==126 for p in PORTALS),
        "kb_parse_zero_fail":all(x["status"]=="PASS" for x in self.manifest),
        "database_stats":self.dbcount>0 and self.schemacount>0 and self.tablecount>0 and self.dbfieldrows>0,
        "three_portal_sources":all(len(p["portal_sources"])==3 for p in self.processes),
        "traceability":bool(self.trace),"gaps_recorded":True}
        for k0,v in checks.items():
            if not v:self.hard.append("DoD:"+k0)
        return {"phase":PHASE,"phase_gate":"PASS" if not self.hard else "FAIL","checks":checks,"hard_failures":sorted(set(self.hard)),
        "kb":{"files":len(self.manifest),"extensions":dict(Counter(x["ext"] for x in self.manifest)),"failures":sum(x["status"]!="PASS" for x in self.manifest)},
        "pages":{"records":len(self.pages),"by_source":self.page_counts},
        "process":{"catalog":len(self.processes),"workbooks":dict(self.bookcounts),"workbooks_total":sum(self.bookcounts.values()),"s0_codes":len(self.s0codes),"s1_codes":len(self.s1codes),
                   "form_rows":self.formrows,"field_rows":self.fieldrows,"field_rows_by_portal":{p:sum(r["portal_code"]==p for r in self.sections["fields"]) for p in PORTALS}},
        "database":{"count":self.dbcount,"schemas":self.schemacount,"tables":self.tablecount,"field_rows":self.dbfieldrows,"relations":self.relrows,"indexes":self.idxrows,"ddl":self.ddl["counts"]},
        "permission_fragments":len(self.perms),"trace_records":len(self.trace),"api_records":len(self.api),
        "gaps":{"count":len(self.gaps),"severity":dict(Counter(g["severity"] for g in self.gaps)),"category":dict(Counter(g["category"] for g in self.gaps))}}
    def write(self,summary):
        if self.contract.exists():shutil.rmtree(self.contract)
        self.contract.mkdir(parents=True)
        jwrite(self.contract/"KB_PARSE_MANIFEST.json",self.manifest);jwrite(self.contract/"summary.json",summary);jwrite(self.contract/"pages.json",self.pages)
        jlwrite(self.contract/"s0_records.jsonl",self.s0);jlwrite(self.contract/"s1_index_records.jsonl",self.s1);jwrite(self.contract/"processes.json",self.processes)
        jlwrite(self.contract/"process_workbook_sheets.jsonl",self.bookmeta);jlwrite(self.contract/"forms.jsonl",self.sections["forms"])
        for p in PORTALS:jlwrite(self.contract/f"fields_{p}.jsonl",(r for r in self.sections["fields"] if r["portal_code"]==p))
        jlwrite(self.contract/"states_approvals.jsonl",self.sections["states"]);jlwrite(self.contract/"rules_interfaces.jsonl",self.sections["rules"]);jlwrite(self.contract/"three_endpoint_linkage.jsonl",self.sections["linkage"])
        jlwrite(self.contract/"permissions.jsonl",self.perms);jlwrite(self.contract/"traceability.jsonl",self.trace);jlwrite(self.contract/"api_records.jsonl",self.api);jwrite(self.contract/"ddl_inventory.json",self.ddl)
        jlwrite(self.contract/"database_process_table_mapping.jsonl",({"process_code":c,**m} for c in sorted(self.pmap) for m in self.pmap[c]))
        self.write_master(summary)
    def write_master(self,s):
        jwrite(self.impl/"MASTER_PAGE_CATALOG.json",{"schema_version":"1.1","generated_in_phase":PHASE,"status":"PARSED","page_count":len(self.pages),"pages":self.pages})
        pg=["# MASTER_PAGE_CATALOG","",f"> 六份页面 Excel 全量实际解析。Page records: **{len(self.pages)}**。","",
            mdtable(["Portal","L1","L2","L3","Display","Process","Scope","Sensitive","Mobile","Source"],[(p["portal_code"],p["level_1"],p["level_2"] or "",p["level_3"] or "",p["display_name"],",".join(p["process_codes"]) or "UNKNOWN",p["data_scope"],p["sensitive_level"],p["mobile_access"],f"{p['source_file']}#{p['source_sheet']}!R{p['source_row']}") for p in self.pages])]
        (self.impl/"MASTER_PAGE_CATALOG.md").write_text("\n".join(pg)+"\n",encoding="utf-8")
        jwrite(self.impl/"MASTER_PROCESS_CATALOG.json",{"schema_version":"1.1","generated_in_phase":PHASE,"status":"PARSED","process_count":len(self.processes),"processes":self.processes})
        pr=["# MASTER_PROCESS_CATALOG","",f"> P001–P126 catalog: **{len(self.processes)}**; three-portal workbooks: **{sum(self.bookcounts.values())}**。","",
            mdtable(["Process","Name","Employee","Center","Tech","DB mapping"],[(p["process_code"],p["canonical_name"],"YES" if "employee" in p["portal_sources"] else "NO","YES" if "center" in p["portal_sources"] else "NO","YES" if "tech" in p["portal_sources"] else "NO","; ".join(f"{m['schema']}.{m['table']}" for m in p["database_mappings"]) or "UNKNOWN") for p in self.processes])]
        (self.impl/"MASTER_PROCESS_CATALOG.md").write_text("\n".join(pr)+"\n",encoding="utf-8")
        jwrite(self.impl/"MASTER_PERMISSION_MATRIX.json",{"generated_in_phase":PHASE,"records":self.perms})
        (self.impl/"MASTER_PERMISSION_MATRIX.md").write_text("# MASTER_PERMISSION_MATRIX\n\n> 仅抽取源资料中可直接来源化的权限/角色/动作/数据范围/敏感/审批字段；无来源保持 UNKNOWN。\n\n"+mdtable(["Portal","Fragments"],[(p,sum(x["portal_code"]==p for x in self.perms)) for p in PORTALS])+"\n\nMachine: `contracts/phase-01/permissions.jsonl`\n",encoding="utf-8")
        dbrows=[(p["process_code"],p["canonical_name"],"; ".join(sorted({m["schema"] for m in p["database_mappings"]})) or "UNKNOWN","; ".join(sorted({m["table"] for m in p["database_mappings"]})) or "UNKNOWN") for p in self.processes]
        (self.impl/"MASTER_DATABASE_MAPPING.md").write_text("# MASTER_DATABASE_MAPPING\n\n"+mdtable(["Metric","Actual","Declared"],[("Database",self.dbcount,3),("Schema",self.schemacount,46),("Table",self.tablecount,265),("Database field rows",self.dbfieldrows,"n/a"),("Relations",self.relrows,"n/a"),("Indexes",self.idxrows,"n/a")])+"\n\n## Process mapping\n\n"+mdtable(["Process","Name","Schema","Table"],dbrows)+"\n",encoding="utf-8")
        jwrite(self.impl/"MASTER_DATABASE_MAPPING.json",{"generated_in_phase":PHASE,"stats":s["database"],"processes":[{"process_code":p["process_code"],"name":p["canonical_name"],"database_mappings":p["database_mappings"]} for p in self.processes]})
        (self.impl/"MASTER_API_CATALOG.md").write_text(f"# MASTER_API_CATALOG\n\n> 只记录 interface catalog 中明确给出的 HTTP method/path；不从流程名猜 API。\n\n- API-like records: **{len(self.api)}**\n- Machine: `contracts/phase-01/api_records.jsonl`\n",encoding="utf-8")
        (self.impl/"MASTER_TRACEABILITY.md").write_text(f"# MASTER_TRACEABILITY\n\n> page(source file/sheet/row) → process → form/field → table → permission → API/interface。UNKNOWN 表示无可来源化事实。\n\n- Pages: **{len(self.pages)}**\n- Processes: **{len(self.processes)}**\n- Trace records: **{len(self.trace)}**\n- Machine: `contracts/phase-01/traceability.jsonl`\n",encoding="utf-8")
        jwrite(self.impl/"MASTER_GAPS.json",self.gaps);(self.impl/"MASTER_GAPS.md").write_text("# MASTER_GAPS\n\n> 只记录资料内部冲突/缺口，不擅自修正业务。\n\n"+(mdtable(["ID","Severity","Category","Source","Description","Impact"],[(g["id"],g["severity"],g["category"],g["source"],g["description"],g["impact"]) for g in self.gaps]) if self.gaps else "无。")+"\n",encoding="utf-8")
        self.worklist();self.report(s);self.progress(s)
    def worklist(self):
        lines=["# PHASE-02–29 PROCESS WORKLIST","","> 基于当前 P001–P126 catalog 按编码顺序生成，只定义后续流程批次，不创造业务规则。后续正式 PHASE 提示词优先。",""];obj=[];cur=0
        for ph,size in zip(range(2,30),[5]*14+[4]*14):
            batch=self.processes[cur:cur+size];cur+=size;lines+= [f"## PHASE-{ph:02d}","",mdtable(["Process","Name"],[(p["process_code"],p["canonical_name"]) for p in batch]),""]
            obj.append({"phase":f"PHASE-{ph:02d}","processes":[{"process_code":p["process_code"],"name":p["canonical_name"],"database_mappings":p["database_mappings"]} for p in batch]})
        (self.impl/"PHASE_02_29_WORKLIST.md").write_text("\n".join(lines)+"\n",encoding="utf-8");jwrite(self.impl/"PHASE_02_29_WORKLIST.json",obj)
    def report(self,s):
        pd=self.impl/"phases/PHASE-01";ev=self.impl/"evidence";pd.mkdir(parents=True,exist_ok=True);ev.mkdir(parents=True,exist_ok=True)
        lines=["# PHASE-01 REPORT","","`PLATFORM/基础工程` — 全量 Knowledge Base 机器化解析与需求追溯。未创建业务页面/API/数据库迁移。","",
        mdtable(["Metric","Actual"],[("KB files",s["kb"]["files"]),("KB parse failures",s["kb"]["failures"]),("Page records",s["pages"]["records"]),("Processes",s["process"]["catalog"]),("Three-portal workbooks",s["process"]["workbooks_total"]),("Form rows",s["process"]["form_rows"]),("Source field rows",s["process"]["field_rows"]),("Databases",s["database"]["count"]),("Schemas",s["database"]["schemas"]),("Tables",s["database"]["tables"]),("DB field rows",s["database"]["field_rows"]),("Gaps",s["gaps"]["count"])]),"","## DoD"]
        lines += [f"- [{'x' if v else ' '}] {k0}" for k0,v in s["checks"].items()];lines += ["",f"## PHASE GATE: {s['phase_gate']}","","GitHub commit/push/PR/CI metadata由控制智能体在观察到远端生成 commit 后补录。"]
        (pd/"PHASE_REPORT.md").write_text("\n".join(lines)+"\n",encoding="utf-8")
        jwrite(ev/"PHASE-01_VALIDATION.json",{"phase":PHASE,"generated_at_utc":dt.datetime.now(dt.timezone.utc).isoformat(),"phase_gate":s["phase_gate"],"required_checks":s["checks"],"hard_failures":s["hard_failures"],"summary":s,"security":{"employee_master_values_exported":False,"sample_redaction":True}})
    def progress(self,s):
        p=self.impl/"MASTER_PROGRESS.md";t=p.read_text(encoding="utf-8")
        t=t.replace("> Current completed phase: PHASE-00",f"> Current completed phase: {PHASE if s['phase_gate']=='PASS' else 'PHASE-00'}")
        t=t.replace("`PHASE-00 = PASS`。PHASE-01 至 PHASE-35 全部保持 `NOT_STARTED`，等待各自阶段提示词；不得自行提前施工。",f"`PHASE-00 = PASS`；`PHASE-01 = {s['phase_gate']}`。PHASE-02 至 PHASE-35 保持 `NOT_STARTED`。")
        t=re.sub(r"\| PHASE-01 \| (?:NOT_STARTED|IN_PROGRESS|PASS|FAIL|BLOCKED) \|[^\n]*",f"| PHASE-01 | {s['phase_gate']} | 全量 Knowledge Base 机器化解析与需求追溯 |",t)
        t += f"\n## PHASE-01 machine snapshot\n\n- KB files: {s['kb']['files']}\n- Pages: {s['pages']['records']}\n- Processes: {s['process']['catalog']}\n- Workbooks: {s['process']['workbooks_total']}\n- Form rows: {s['process']['form_rows']}\n- Source field rows: {s['process']['field_rows']}\n- DB/Schema/Table: {s['database']['count']}/{s['database']['schemas']}/{s['database']['tables']}\n- Gaps: {s['gaps']['count']}\n- PHASE-01: {s['phase_gate']}\n"
        p.write_text(t,encoding="utf-8")
    def run(self):
        if not self.kb.exists():raise SystemExit("Knowledge Base missing")
        self.scan();self.parse_pages();self.parse_indexes_and_processes();self.parse_db();self.build_process_catalog();self.build_trace()
        if not (self.kb/"04 Agents开发规范").exists():self.gap("INFO","KB_STRUCTURE","Knowledge Base/","04 Agents开发规范 当前不存在；根 AGENT 仍为 canonical","不自行补造第二份规则")
        s=self.summary();self.write(s);print(json.dumps(s,ensure_ascii=False,indent=2));return 0 if s["phase_gate"]=="PASS" else 3

def main():
    ap=argparse.ArgumentParser();ap.add_argument("--repo-root",default=".");a=ap.parse_args();return Pipeline(a.repo_root).run()
if __name__=="__main__":raise SystemExit(main())
