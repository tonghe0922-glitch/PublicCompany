#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

PAGE_FILES = [
    "Knowledge Base/01 完整的页面架构/1-1 员工首页.xlsx",
    "Knowledge Base/01 完整的页面架构/1-2员工全层级页面.xlsx",
    "Knowledge Base/01 完整的页面架构/2-1 中心首页.xlsx",
    "Knowledge Base/01 完整的页面架构/2-2中心全层级页面.xlsx",
    "Knowledge Base/01 完整的页面架构/3-1技术-首页.xlsx",
    "Knowledge Base/01 完整的页面架构/3-2技术-全层级页面.xlsx",
]
S0 = "Knowledge Base/02 业务流程 表单 字段/S0 全部业务流程简表.xlsx"
S1 = "Knowledge Base/02 业务流程 表单 字段/S1 三端业务流程表单字段包/00_三端业务流程表单字段总索引.xlsx"
REPRESENTATIVE = [
    "Knowledge Base/02 业务流程 表单 字段/S1 三端业务流程表单字段包/01_员工端/01_平台公共能力/001_统一登录与多岗位身份切换.xlsx",
    "Knowledge Base/02 业务流程 表单 字段/S1 三端业务流程表单字段包/02_中心管理端/01_平台公共能力/001_统一登录与多岗位身份切换.xlsx",
    "Knowledge Base/02 业务流程 表单 字段/S1 三端业务流程表单字段包/03_技术后台端/01_平台公共能力/001_统一登录与多岗位身份切换.xlsx",
    "Knowledge Base/02 业务流程 表单 字段/S1 三端业务流程表单字段包/01_员工端/03_财人中心/025_报销闭环.xlsx",
]
CSV_FILES = [
    "Knowledge Base/03 数据库需求规则/02_数据字典/03_全量表清单.csv",
    "Knowledge Base/03 数据库需求规则/02_数据字典/04_全量字段字典.csv",
    "Knowledge Base/03 数据库需求规则/02_数据字典/07_流程落表映射.csv",
    "Knowledge Base/03 数据库需求规则/04_初始化数据/01_process_catalog.csv",
    "Knowledge Base/03 数据库需求规则/04_初始化数据/04_interface_catalog.csv",
]


def txt(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).replace("\r", "\n").split())


def safe(value: Any) -> Any:
    value = txt(value)
    if not value:
        return ""
    return value if len(value) <= 220 else value[:217] + "..."


def workbook_sample(path: Path, limit_rows: int = 14) -> dict[str, Any]:
    wb = load_workbook(path, read_only=True, data_only=False)
    result: dict[str, Any] = {"path": path.as_posix(), "sheets": []}
    try:
        for ws in wb.worksheets:
            rows = []
            for row_number, row in enumerate(ws.iter_rows(values_only=True), 1):
                values = [safe(value) for value in row]
                while values and values[-1] == "":
                    values.pop()
                if not any(values):
                    continue
                rows.append({"row": row_number, "values": values})
                if len(rows) >= limit_rows:
                    break
            result["sheets"].append({
                "title": ws.title,
                "max_row": ws.max_row,
                "max_column": ws.max_column,
                "first_nonempty_rows": rows,
            })
    finally:
        wb.close()
    return result


def read_csv_sample(path: Path, limit_rows: int = 8) -> dict[str, Any]:
    raw = path.read_bytes()
    decoded = None
    encoding = None
    for candidate in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            decoded = raw.decode(candidate)
            encoding = candidate
            break
        except UnicodeDecodeError:
            pass
    if decoded is None:
        return {"path": path.as_posix(), "error": "unsupported encoding"}

    try:
        dialect = csv.Sniffer().sniff(decoded[:65536], delimiters=",\t;")
    except csv.Error:
        dialect = csv.excel
    rows = []
    for row_number, row in enumerate(csv.reader(decoded.splitlines(), dialect), 1):
        if not any(txt(value) for value in row):
            continue
        rows.append({"row": row_number, "values": [safe(value) for value in row]})
        if len(rows) >= limit_rows:
            break
    return {"path": path.as_posix(), "encoding": encoding, "rows": rows}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--repo-root", default=".")
    args = parser.parse_args()
    root = Path(args.repo_root).resolve()
    output = root / "docs/implementation/contracts/phase-01/diagnostics"
    output.mkdir(parents=True, exist_ok=True)

    payloads = {
        "page_workbook_samples.json": [workbook_sample(root / path) for path in PAGE_FILES],
        "process_index_samples.json": [workbook_sample(root / S0, 20), workbook_sample(root / S1, 20)],
        "process_workbook_samples.json": [workbook_sample(root / path, 10) for path in REPRESENTATIVE],
        "database_csv_samples.json": [read_csv_sample(root / path, 10) for path in CSV_FILES],
    }
    for filename, payload in payloads.items():
        (output / filename).write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Wrote parser diagnostics to {output.relative_to(root).as_posix()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
